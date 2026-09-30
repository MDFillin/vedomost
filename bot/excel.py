"""Ведомость в Excel — в формате бумажного журнала старосты.

Лист на каждый месяц:
  строка 1: название группы | дни недели
  строка 2: «МЕСЯЦ ГОД»      | числа месяца | ИТОГО
  далее:    ФИО студента     | «н» — не был, «б» — болел, пусто — был | число «н»
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .db import Database, Session
from .utils import MONTHS, WEEKDAYS_SHORT, to_local

FONT = "Times New Roman"
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW = PatternFill("solid", fgColor="FFFF00")
SICK_FILL = PatternFill("solid", fgColor="FBE4D5")  # светло-оранжевый, как в журнале
CENTER = Alignment(horizontal="center", vertical="center")


def month_key(iso_date: str) -> str:
    return iso_date[:7]


def month_title(ym: str) -> str:
    y, m = map(int, ym.split("-"))
    return f"{MONTHS[m - 1]} {y}"


def _cell(ws, row: int, col: int, value=None, *, bold=False, color=None, fill=None,
          align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=11, bold=bold, color=color)
    c.border = BORDER
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def _month_sheet(wb: Workbook, db: Database, ym: str, sessions: list[Session],
                 group_name: str) -> None:
    ws = wb.create_sheet(month_title(ym)[:31])
    students = sorted(db.list_students(), key=lambda s: s.name.lower())
    leaves = db.sick_days_by_user()
    sessions = sorted(sessions, key=lambda s: (s.date, s.time or ""))
    per_day: dict[str, int] = {}
    for s in sessions:
        per_day[s.date] = per_day.get(s.date, 0) + 1

    _cell(ws, 1, 1, group_name, bold=True)
    _cell(ws, 2, 1, f" {month_title(ym).upper()}", bold=True, fill=YELLOW,
          align=Alignment(vertical="center"))
    ws.column_dimensions["A"].width = 37

    for i, s in enumerate(sessions):
        col = i + 2
        d = date.fromisoformat(s.date)
        header = WEEKDAYS_SHORT[d.weekday()]
        many = per_day[s.date] > 1 and s.time
        if many:  # несколько пар в день — подписываем время
            header += f" {s.time}"
        _cell(ws, 1, col, header, align=CENTER if many else None)
        _cell(ws, 2, col, d.day, align=CENTER)
        ws.column_dimensions[get_column_letter(col)].width = 9 if many else 4

    total_col = len(sessions) + 2
    last = get_column_letter(total_col - 1)
    _cell(ws, 1, total_col)
    _cell(ws, 2, total_col, "ИТОГО", bold=True, color="FF0000", align=CENTER)
    ws.column_dimensions[get_column_letter(total_col)].width = 9

    marks = {s.id: db.marks(s.id) for s in sessions}
    for r, st in enumerate(students, start=3):
        _cell(ws, r, 1, st.name, align=Alignment(indent=1))
        for i, s in enumerate(sessions):
            value, fill = None, None
            if st.user_id not in marks[s.id]:
                if any(lv.covers(s.date) for lv in leaves.get(st.user_id, [])):
                    value, fill = "б", SICK_FILL
                else:
                    value = "н"
            _cell(ws, r, i + 2, value, fill=fill, align=CENTER)
        formula = f'=COUNTIFS(B{r}:{last}{r},"Н")' if sessions else 0
        _cell(ws, r, total_col, formula, bold=True, align=CENTER)

    ws.freeze_panes = "B3"
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _sick_sheet(wb: Workbook, db: Database) -> None:
    leaves = db.list_sick_leaves()
    if not leaves:
        return
    ws = wb.create_sheet("Больничные")
    headers = ["Студент", "С", "По", "Дней", "Кто указал", "Записано"]
    widths = [37, 12, 12, 7, 14, 17]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        _cell(ws, 1, i, h, bold=True, fill=YELLOW, align=CENTER)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, lv in enumerate(sorted(leaves, key=lambda x: (x.start, x.user_id)), start=2):
        st = db.get_student(lv.user_id)
        _cell(ws, r, 1, st.name if st else str(lv.user_id), align=Alignment(indent=1))
        # настоящие даты, чтобы работала формула «Дней»
        _cell(ws, r, 2, date.fromisoformat(lv.start), align=CENTER).number_format = "DD.MM.YYYY"
        _cell(ws, r, 3, date.fromisoformat(lv.end), align=CENTER).number_format = "DD.MM.YYYY"
        _cell(ws, r, 4, f"=C{r}-B{r}+1", align=CENTER)
        _cell(ws, r, 5, "староста" if lv.by_admin else "студент", align=CENTER)
        _cell(ws, r, 6, f"{to_local(db, lv.created_at):%d.%m.%Y %H:%M}", align=CENTER)
    ws.freeze_panes = "A2"


def build_workbook(db: Database, months: list[str] | None = None) -> bytes:
    """Ведомость за выбранные месяцы (YYYY-MM); None — за все месяцы."""
    group_name = db.get("group_name") or "Группа"
    by_month: dict[str, list[Session]] = {}
    for s in db.list_sessions():
        by_month.setdefault(month_key(s.date), []).append(s)

    wb = Workbook()
    wb.remove(wb.active)
    selected = sorted(by_month) if months is None else sorted(m for m in months if m in by_month)
    for ym in selected:
        _month_sheet(wb, db, ym, by_month[ym], group_name)
    if not selected:
        ws = wb.create_sheet("Список")
        _cell(ws, 1, 1, group_name, bold=True)
        _cell(ws, 2, 1, "Занятий пока не было")
    _sick_sheet(wb, db)
    wb.calculation.fullCalcOnLoad = True  # Excel пересчитает ИТОГО при открытии

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
