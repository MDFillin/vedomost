"""Тесты без сети: Telegram API подменяется фейковой сессией."""

import asyncio
import io
from datetime import date, datetime, timedelta
from typing import Any

import openpyxl
import pytest
from aiogram import Bot, Dispatcher
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.base import BaseSession
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.methods import (
    AnswerCallbackQuery,
    EditMessageText,
    GetChatAdministrators,
    GetChatMember,
    GetChatMemberCount,
    GetMe,
    SendDocument,
    SendMessage,
    TelegramMethod,
)
from aiogram.types import (
    CallbackQuery,
    Chat,
    ChatMemberLeft,
    ChatMemberMember,
    ChatMemberOwner,
    ChatMemberUpdated,
    Message,
    Update,
    User,
)

from bot.config import Config
from bot.db import Database
from bot.handlers import admin, errors, group, sick
from bot import members
from bot.service import compute_close_at, due_slots, next_slot, send_session
from bot.utils import get_tz, parse_date, parse_times

ADMIN = 100
GROUP = -1001
OUTSIDER = 999  # не состоит в чате группы
STUDENTS = [(201, "Иван Иванов"), (202, "Пётр Петров"), (203, "Анна Смирнова")]


class FakeSession(BaseSession):
    def __init__(self) -> None:
        super().__init__()
        self.calls: list[TelegramMethod] = []
        self.next_id = 1000

    async def make_request(self, bot: Bot, method: TelegramMethod, timeout: Any = None) -> Any:
        self.calls.append(method)
        if isinstance(method, (SendMessage, SendDocument)):
            self.next_id += 1
            return Message(
                message_id=self.next_id,
                date=datetime.now(),
                chat=Chat(id=method.chat_id, type="private" if method.chat_id > 0 else "supergroup"),
                text=getattr(method, "text", None),
            )
        if isinstance(method, EditMessageText):
            return True
        if isinstance(method, GetChatMember):
            u = User(id=method.user_id, is_bot=False, first_name="X")
            if method.user_id == OUTSIDER:
                return ChatMemberLeft(user=u)
            return ChatMemberMember(user=u)
        if isinstance(method, GetChatAdministrators):
            return [ChatMemberOwner(user=user(ADMIN), is_anonymous=False)]
        if isinstance(method, GetChatMemberCount):
            return 31
        if isinstance(method, GetMe):
            return User(id=42, is_bot=True, first_name="Ведомость", username="vedomost_bot")
        return True

    async def close(self) -> None:
        pass

    async def stream_content(self, *a: Any, **kw: Any):  # pragma: no cover
        yield b""

    def of(self, kind: type) -> list:
        return [c for c in self.calls if isinstance(c, kind)]


@pytest.fixture
def env():
    db = Database(":memory:", "Europe/Moscow")
    config = Config(token="1:x", admin_ids=frozenset({ADMIN}), db_path=":memory:",
                    default_tz="Europe/Moscow")
    session = FakeSession()
    bot = Bot("42:TEST", session=session, default=DefaultBotProperties(parse_mode="HTML"))
    dp = Dispatcher(storage=MemoryStorage(), db=db, config=config)
    routers = [group.router, sick.router, *admin.setup(config)]
    for r in routers:
        dp.include_router(r)
    dp.errors.register(errors.on_error)
    members._member_cache.clear()
    yield db, bot, dp, session
    for r in routers:  # роутеры модульные — отцепляем для следующего теста
        r._parent_router = None


_uid = iter(range(1, 10**6))


def user(uid: int, name: str = "Староста") -> User:
    first, *last = name.split()
    return User(id=uid, is_bot=False, first_name=first, last_name=" ".join(last) or None,
                username=f"user{uid}")


async def text(dp, bot, uid: int, txt: str, chat_id: int | None = None, chat_type="private"):
    chat_id = chat_id or uid
    msg = Message(message_id=next(_uid), date=datetime.now(),
                  chat=Chat(id=chat_id, type=chat_type), from_user=user(uid), text=txt)
    await dp.feed_update(bot, Update(update_id=next(_uid), message=msg))


async def press(dp, bot, uid: int, data: str, chat_id: int, name: str = "Староста"):
    msg = Message(message_id=next(_uid), date=datetime.now(),
                  chat=Chat(id=chat_id, type="private" if chat_id > 0 else "supergroup"),
                  text="…")
    cb = CallbackQuery(id=str(next(_uid)), from_user=user(uid, name), chat_instance="x",
                       message=msg, data=data)
    await dp.feed_update(bot, Update(update_id=next(_uid), callback_query=cb))


def last_alert(session: FakeSession) -> str:
    return session.of(AnswerCallbackQuery)[-1].text or ""


def buttons(method) -> list[str]:
    markup = method.reply_markup
    if not markup:
        return []
    return [b.callback_data for row in markup.inline_keyboard for b in row if b.callback_data]


# ---------- юнит-тесты ----------

def test_parsers():
    today = date(2026, 9, 30)
    assert parse_times("9:00 13.40") == ["09:00", "13:40"]
    assert parse_times("25:00") is None
    assert parse_date("15.10", today) == date(2026, 10, 15)
    assert parse_date("01.09", today) == date(2027, 9, 1)
    assert parse_date("04.11.26", today) == date(2026, 11, 4)
    assert parse_date("31.02.2027", today) is None


def test_schedule_due_and_skip():
    db = Database(":memory:", "Europe/Moscow")
    tz = get_tz(db)
    db.add_schedule(2, "10:00")  # среда
    wed = datetime(2026, 9, 30, 10, 5, tzinfo=tz)
    assert due_slots(db, wed) == [(date(2026, 9, 30), "10:00")]
    assert due_slots(db, wed + timedelta(minutes=20)) == []  # опоздание больше окна
    assert due_slots(db, wed - timedelta(minutes=10)) == []
    db.add_skip_date("2026-09-30")
    assert due_slots(db, wed) == []
    db.add_extra_date("2026-09-30", "10:00")  # разовая дата важнее пропуска
    assert due_slots(db, wed) == [(date(2026, 9, 30), "10:00")]
    assert next_slot(db, wed) == datetime(2026, 10, 7, 10, 0, tzinfo=tz)


def test_mark_is_atomic():
    db = Database(":memory:")
    s = db.create_session(GROUP, "2026-09-30", "10:00", "slot", None)
    assert db.mark(s.id, 1) is True
    assert db.mark(s.id, 1) is False
    assert db.create_session(GROUP, "2026-09-30", "10:00", "slot", None) is None


# ---------- сценарии ----------

async def test_full_flow(env):
    db, bot, dp, session = env

    # не-админ не может привязать группу
    await text(dp, bot, 201, "/bind", GROUP, "supergroup")
    assert db.group_chat_id is None
    await text(dp, bot, ADMIN, "/bind", GROUP, "supergroup")
    assert db.group_chat_id == GROUP

    # регистрация
    s = await send_session(bot, db, date.today(), "10:00", slot="x")
    group_msg = session.of(SendMessage)[-1]
    assert group_msg.chat_id == GROUP
    assert buttons(group_msg) == [f"mark:{s.id}", f"sick:{s.id}"]
    assert group_msg.reply_markup.inline_keyboard[0][0].style == "success"
    for uid, name in STUDENTS:
        await press(dp, bot, uid, "register", GROUP, name)
    assert len(db.list_students()) == 4  # 3 студента + староста (админ чата)

    # отметка + повторное нажатие
    await press(dp, bot, 201, f"mark:{s.id}", GROUP, "Иван Иванов")
    assert "записано" in last_alert(session)
    await press(dp, bot, 201, f"mark:{s.id}", GROUP, "Иван Иванов")
    assert "уже отметились" in last_alert(session)
    await press(dp, bot, 202, f"mark:{s.id}", GROUP, "Пётр Петров")
    assert set(db.marks(s.id)) == {201, 202}

    # счётчик в группе обновляется с задержкой
    await asyncio.sleep(2.2)
    edits = [e for e in session.of(EditMessageText) if e.chat_id == GROUP]
    assert edits and "Отметились: <b>2</b>" in edits[-1].text

    # карточка занятия у старосты
    await press(dp, bot, ADMIN, f"a:s:{s.id}", ADMIN)
    card = session.of(EditMessageText)[-1].text
    assert "Были: 2/4" in card and "Анна Смирнова" in card.split("Не было")[1]

    # закрытие — нажать больше нельзя
    await press(dp, bot, ADMIN, f"a:sc:{s.id}", ADMIN)
    await press(dp, bot, 203, f"mark:{s.id}", GROUP, "Анна Смирнова")
    assert "закрыта" in last_alert(session)

    # староста отмечает вручную
    await press(dp, bot, ADMIN, f"a:st:{s.id}:203:0", ADMIN)
    assert 203 in db.marks(s.id) and db.marks(s.id)[203].by_admin

    # не-админ в личке панель не видит
    before = len(session.of(SendMessage))
    await text(dp, bot, 201, "/menu")
    assert "Отмечайтесь кнопкой" in session.of(SendMessage)[before].text
    await press(dp, bot, 201, "a:menu", 201)
    assert len(session.of(SendMessage)) == before + 1  # кнопка проигнорирована


async def test_schedule_input(env):
    db, bot, dp, session = env
    await press(dp, bot, ADMIN, "a:schw", ADMIN)
    await press(dp, bot, ADMIN, "a:schwt:0", ADMIN)
    await press(dp, bot, ADMIN, "a:schwt:3", ADMIN)
    await press(dp, bot, ADMIN, "a:schwn", ADMIN)
    await text(dp, bot, ADMIN, "9:00 13:40")
    assert [(wd, t) for _, wd, t in db.list_schedule()] == [
        (0, "09:00"), (0, "13:40"), (3, "09:00"), (3, "13:40")]

    await press(dp, bot, ADMIN, "a:schd", ADMIN)
    await text(dp, bot, ADMIN, "15.10.2099 10:30\n16.10.2099 11:00")
    assert len(db.list_extra_dates()) == 2

    await press(dp, bot, ADMIN, "a:schs", ADMIN)
    await text(dp, bot, ADMIN, "29.12.2099-02.01.2100")
    assert len(db.list_skip_dates()) == 5

    await press(dp, bot, ADMIN, "a:schx", ADMIN)
    items = [b for b in buttons(session.of(EditMessageText)[-1]) if b.startswith("a:schx")]
    for data in items:
        await press(dp, bot, ADMIN, data, ADMIN)
    assert not db.list_schedule() and not db.list_extra_dates() and not db.list_skip_dates()


async def test_settings_input(env):
    db, bot, dp, session = env
    await press(dp, bot, ADMIN, "a:seti:timezone", ADMIN)
    await text(dp, bot, ADMIN, "Mars/Olympus")
    assert db.get("timezone") == "Europe/Moscow"
    await text(dp, bot, ADMIN, "Asia/Novosibirsk")
    assert db.get("timezone") == "Asia/Novosibirsk"
    await press(dp, bot, ADMIN, "a:seti:button", ADMIN)
    await text(dp, bot, ADMIN, "✅ Я на паре")
    assert db.get("button_text") == "✅ Я на паре"
    await press(dp, bot, ADMIN, "a:setwv:0", ADMIN)
    assert db.get_int("window_minutes") == 0


async def test_click_every_admin_button(env):
    """Обходит все кнопки панели и проверяет, что ни одна не падает."""
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    for uid, name in STUDENTS:
        db.upsert_student(uid, name, None)
    db.add_schedule(0, "09:00")
    db.add_extra_date("2099-01-01", "10:00")
    db.add_skip_date("2099-01-02")
    s = await send_session(bot, db, date.today(), "10:00")
    db.mark(s.id, 201)
    db.rename_student(202, "Петров П.")

    skip = {"a:sdy:", "a:udy:", "a:schxsa", "a:schxw", "a:schxd", "a:schxs:"}
    seen: set[str] = set()
    queue = ["a:menu"]
    errors = []

    dp.errors.handlers.clear()  # в этом тесте ошибки не глушим, а собираем
    dp.errors.register(lambda event: errors.append(event.exception))
    while queue:
        data = queue.pop()
        if data in seen or any(data.startswith(p) for p in skip):
            continue
        seen.add(data)
        n = len(session.calls)
        await press(dp, bot, ADMIN, data, ADMIN)
        for call in session.calls[n:]:
            if isinstance(call, (EditMessageText, SendMessage)) and call.chat_id == ADMIN:
                queue += [b for b in buttons(call) if b.startswith("a:")]
        assert session.of(AnswerCallbackQuery), data

    assert not errors, errors
    assert len(seen) > 40
    # удаление в самом конце
    await press(dp, bot, ADMIN, f"a:sdy:{s.id}", ADMIN)
    await press(dp, bot, ADMIN, "a:udy:201", ADMIN)
    assert db.get_session(s.id) is None and db.get_student(201) is None
    assert not errors, errors


# ---------- новые функции ----------

async def test_protection(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    s = await send_session(bot, db, date.today(), "10:00")
    assert session.of(SendMessage)[-1].protect_content is True  # запрет пересылки

    # посторонний, не состоящий в чате
    await press(dp, bot, OUTSIDER, f"mark:{s.id}", GROUP, "Чужой Человек")
    assert "только участники" in last_alert(session)
    assert db.get_student(OUTSIDER) is None  # в список не попал

    # нажатие с пересланной копии в другом чате
    await press(dp, bot, 201, f"mark:{s.id}", -5555, "Иван Иванов")
    assert "только в чате группы" in last_alert(session)
    await press(dp, bot, 201, f"mark:{s.id}", 201, "Иван Иванов")
    assert "только в чате группы" in last_alert(session)
    assert not db.marks(s.id)

    # исключённый старостой
    db.upsert_student(202, "Пётр Петров", None)
    db.set_student_active(202, False)
    await press(dp, bot, 202, f"mark:{s.id}", GROUP, "Пётр Петров")
    assert "исключил" in last_alert(session)

    # «только из списка»: новичок не может, студент из списка — может
    db.set("roster_only", "1")
    await press(dp, bot, 203, f"mark:{s.id}", GROUP, "Анна Смирнова")
    assert "нет в списке" in last_alert(session)
    db.upsert_student(203, "Анна Смирнова", None)
    await press(dp, bot, 203, f"mark:{s.id}", GROUP, "Анна Смирнова")
    assert "записано" in last_alert(session)

    # при выключенной проверке членства посторонний пройдёт
    db.set("roster_only", "0")
    db.set("members_only", "0")
    await press(dp, bot, OUTSIDER, f"mark:{s.id}", GROUP, "Чужой Человек")
    assert "записано" in last_alert(session)


async def test_day_mode(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    db.set("mode", "day")
    first = await send_session(bot, db, date.today(), "09:00", slot="a")
    assert first is not None and first.time is None
    assert await send_session(bot, db, date.today(), "13:40", slot="b") is None
    await press(dp, bot, ADMIN, "a:sendy", ADMIN)
    assert "уже" in last_alert(session)
    assert db.count_sessions() == 1

    db.set("mode", "pair")
    assert await send_session(bot, db, date.today(), "13:40", slot="c") is not None


def test_close_modes():
    db = Database(":memory:", "Europe/Moscow")
    tz = get_tz(db)
    start = datetime(2026, 9, 30, 9, 0, tzinfo=tz)
    db.set("window_minutes", 45)
    assert compute_close_at(db, start) == start + timedelta(minutes=45)
    db.set("close_mode", "until")
    db.set("close_until", "10:30")
    assert compute_close_at(db, start) == datetime(2026, 9, 30, 10, 30, tzinfo=tz)
    late = datetime(2026, 9, 30, 11, 0, tzinfo=tz)  # отправили позже 10:30 → до конца дня
    assert compute_close_at(db, late) == datetime(2026, 9, 30, 23, 59, tzinfo=tz)
    db.set("close_mode", "eod")
    assert compute_close_at(db, start) == datetime(2026, 9, 30, 23, 59, tzinfo=tz)
    db.set("close_mode", "none")
    assert compute_close_at(db, start) is None


async def test_close_setting_applies_to_open_session(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    s = await send_session(bot, db, date.today(), "10:00")
    await press(dp, bot, ADMIN, "a:setwm:none", ADMIN)
    assert db.get_session(s.id).close_at is None
    await press(dp, bot, ADMIN, "a:setwu", ADMIN)
    await text(dp, bot, ADMIN, "23:58")
    assert db.get("close_mode") == "until"
    assert to_hm(db, db.get_session(s.id).close_at) in {"23:58", "23:59"}


def to_hm(db, iso):
    from bot.utils import to_local
    return f"{to_local(db, iso):%H:%M}"


async def test_message_template(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    await press(dp, bot, ADMIN, "a:seti:message_text", ADMIN)
    await text(dp, bot, ADMIN, "Пара {время}, {день} {дата}. Успей до {до}!")
    await press(dp, bot, ADMIN, "a:seti:title", ADMIN)
    await text(dp, bot, ADMIN, "ИВТ-21 <тест>")
    db.set("close_mode", "until")
    db.set("close_until", "23:59")
    await send_session(bot, db, date(2026, 9, 30), "10:00")
    msg = session.of(SendMessage)[-1].text
    assert "Пара 10:00, Среда 30.09.2026. Успей до 23:59!" in msg
    assert "ИВТ-21 &lt;тест&gt;" in msg  # текст экранируется, HTML не ломается

    await press(dp, bot, ADMIN, "a:setpv", ADMIN)
    assert "Успей до" in session.of(SendMessage)[-1].text
    await press(dp, bot, ADMIN, "a:setmr", ADMIN)
    assert db.get("message_text").startswith("Если вы на занятии")


def member_update(chat_id: int, by: int, old: str, new: str, who: User) -> ChatMemberUpdated:
    cls = {"member": ChatMemberMember, "left": ChatMemberLeft}
    return ChatMemberUpdated(
        chat=Chat(id=chat_id, type="supergroup", title="ИВТ-21"),
        from_user=user(by), date=datetime.now(),
        old_chat_member=cls[old](user=who), new_chat_member=cls[new](user=who),
    )


async def test_bot_added_auto_binds_and_reads_members(env, monkeypatch):
    db, bot, dp, session = env
    object.__setattr__(dp["config"], "api_id", 1)
    object.__setattr__(dp["config"], "api_hash", "x")

    async def fake_fetch(config, chat_id):
        return [(ADMIN, "Староста", None), *[(uid, n, None) for uid, n in STUDENTS]]

    monkeypatch.setattr(members, "fetch_all_members", fake_fetch)
    me = User(id=42, is_bot=True, first_name="bot")

    # добавил посторонний — не привязываемся
    await dp.feed_update(bot, Update(update_id=next(_uid),
                                     my_chat_member=member_update(-7, 555, "left", "member", me)))
    assert db.group_chat_id is None

    # добавил староста — привязка и считывание всех
    await dp.feed_update(bot, Update(update_id=next(_uid),
                                     my_chat_member=member_update(GROUP, ADMIN, "left", "member", me)))
    assert db.group_chat_id == GROUP
    assert len(db.list_students()) == 4
    report = [m.text for m in session.of(SendMessage) if m.chat_id == ADMIN][-1]
    assert "Отметки будут приходить в «ИВТ-21»" in report and "<b>4</b>" in report

    # кто-то вышел — после повторной сверки он не считается
    async def fake_fetch2(config, chat_id):
        return [(ADMIN, "Староста", None), *[(uid, n, None) for uid, n in STUDENTS[:2]]]

    monkeypatch.setattr(members, "fetch_all_members", fake_fetch2)
    await press(dp, bot, ADMIN, "a:sync", ADMIN)
    assert 203 not in {s.user_id for s in db.list_students()}
    assert db.get_student(203).in_chat is False


async def test_member_join_leave_without_api(env):
    db, bot, dp, session = env
    await text(dp, bot, ADMIN, "/bind", GROUP, "supergroup")
    assert db.get_student(ADMIN) is not None  # админы чата добавлены через Bot API
    report = [m.text for m in session.of(SendMessage) if m.chat_id == ADMIN][-1]
    assert "API_ID" in report

    newbie = user(301, "Новый Студент")
    msg = Message(message_id=next(_uid), date=datetime.now(),
                  chat=Chat(id=GROUP, type="supergroup"), from_user=newbie,
                  new_chat_members=[newbie])
    await dp.feed_update(bot, Update(update_id=next(_uid), message=msg))
    assert db.get_student(301).counted

    await text(dp, bot, 302, "всем привет", GROUP, "supergroup")
    assert db.get_student(302).counted

    msg = Message(message_id=next(_uid), date=datetime.now(),
                  chat=Chat(id=GROUP, type="supergroup"), from_user=newbie,
                  left_chat_member=newbie)
    await dp.feed_update(bot, Update(update_id=next(_uid), message=msg))
    assert not db.get_student(301).counted


# ---------- больничный ----------

async def test_sick_leave_flow(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    for uid, name in STUDENTS:
        db.upsert_student(uid, name, None)
    today = date.today()
    s = await send_session(bot, db, today, "10:00")

    # кнопка в группе открывает личку с ботом
    await press(dp, bot, 202, f"sick:{s.id}", GROUP, "Пётр Петров")
    assert session.of(AnswerCallbackQuery)[-1].url == "https://t.me/vedomost_bot?start=sick"
    # посторонний не может
    await press(dp, bot, OUTSIDER, f"sick:{s.id}", GROUP, "Чужой Человек")
    assert "только участники" in last_alert(session)

    # в личке: /start sick → календарь → первый и последний день → сохранить
    await text(dp, bot, 202, "/start sick")
    cal = session.of(SendMessage)[-1]
    assert "первый день" in cal.text
    assert f"sk:d:{today.isoformat()}" in buttons(cal)
    start, end = today - timedelta(days=1), today + timedelta(days=2)
    await press(dp, bot, 202, f"sk:d:{start.isoformat()}", 202, "Пётр Петров")
    step2 = session.of(EditMessageText)[-1]
    assert "последний день" in step2.text
    assert f"sk:d:{(start - timedelta(days=1)).isoformat()}" not in buttons(step2)
    await press(dp, bot, 202, f"sk:d:{end.isoformat()}", 202, "Пётр Петров")
    assert "(4 дн.)" in session.of(EditMessageText)[-1].text
    await press(dp, bot, 202, "sk:save", 202, "Пётр Петров")
    assert "записан" in session.of(EditMessageText)[-1].text
    [leave] = db.list_sick_leaves(202)
    assert (leave.start, leave.end, leave.by_admin) == (start.isoformat(), end.isoformat(), False)

    # старосте пришло уведомление, а в группе — счётчик болеющих
    assert any("на больничном" in m.text for m in session.of(SendMessage) if m.chat_id == ADMIN)
    assert "На больничном: <b>1</b>" in [e for e in session.of(EditMessageText)
                                         if e.chat_id == GROUP][-1].text

    # в карточке занятия: отдельный список болеющих
    await press(dp, bot, ADMIN, f"a:s:{s.id}", ADMIN)
    card = session.of(EditMessageText)[-1].text
    assert "На больничном: 1" in card and "Не было без причины: 2" in card

    # итоговая ведомость в Excel
    await press(dp, bot, ADMIN, f"a:xl:{today:%Y-%m}", ADMIN)
    wb = openpyxl.load_workbook(io.BytesIO(session.of(SendDocument)[-1].document.data))
    ws = wb.worksheets[0]
    rows = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(2, ws.max_column + 1)]
            for r in range(3, ws.max_row + 1)}
    assert rows["Пётр Петров"][0] == "б"
    assert rows["Иван Иванов"][0] == "н"
    assert rows["Иван Иванов"][-1] == '=COUNTIFS(B4:B4,"Н")'
    assert "Больничные" in wb.sheetnames

    # студент видит свои больничные и может удалить
    await press(dp, bot, 202, "sk:my", 202, "Пётр Петров")
    assert f"sk:del:{leave.id}" in buttons(session.of(EditMessageText)[-1])
    await press(dp, bot, 202, f"sk:del:{leave.id}", 202, "Пётр Петров")
    assert not db.list_sick_leaves(202)


async def test_sick_leave_typed_and_by_admin(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    db.upsert_student(201, "Иван Иванов", None)
    today = date.today()

    # даты текстом
    await press(dp, bot, 201, "sk:new", 201, "Иван Иванов")
    a, b = today - timedelta(days=3), today
    await text(dp, bot, 201, f"с {a:%d.%m} по {b:%d.%m}")
    assert "(4 дн.)" in session.of(SendMessage)[-1].text
    await press(dp, bot, 201, "sk:save", 201, "Иван Иванов")
    assert db.list_sick_leaves(201)[0].start == a.isoformat()

    # чужой больничный удалить нельзя
    await press(dp, bot, 203, f"sk:del:{db.list_sick_leaves(201)[0].id}", 203, "Анна")
    assert db.list_sick_leaves(201)

    # староста добавляет больничный за студента из его карточки
    await press(dp, bot, ADMIN, "sk:for:201", ADMIN)
    await press(dp, bot, ADMIN, f"sk:d:{today.isoformat()}", ADMIN)
    await press(dp, bot, ADMIN, f"sk:d:{today.isoformat()}", ADMIN)
    await press(dp, bot, ADMIN, "sk:save", ADMIN)
    assert any(lv.by_admin for lv in db.list_sick_leaves(201))

    # раздел «Больничные» у старосты и удаление
    await press(dp, bot, ADMIN, "a:sick", ADMIN)
    dels = [d for d in buttons(session.of(EditMessageText)[-1]) if d.startswith("a:sickdel:")]
    assert len(dels) == 2
    for d in dels:
        await press(dp, bot, ADMIN, d, ADMIN)
    assert not db.list_sick_leaves()

    # не-старосте нельзя добавлять за других
    await press(dp, bot, 203, "sk:for:201", 203, "Анна")
    assert not db.list_sick_leaves()

    # кнопку можно выключить
    await press(dp, bot, ADMIN, "a:sett:sick_button", ADMIN)
    s = await send_session(bot, db, today, "12:00")
    assert buttons(session.of(SendMessage)[-1]) == [f"mark:{s.id}"]


async def test_student_start_menu(env):
    db, bot, dp, session = env
    await text(dp, bot, 201, "/start")
    assert buttons(session.of(SendMessage)[-1]) == ["sk:new", "sk:my"]


# ---------- месяцы, ведомость, дата ----------

async def test_months_and_excel_format(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    db.set("group_name", "УФРС24-2")
    for uid, name in STUDENTS:
        db.upsert_student(uid, name, None)
    sep = [await send_session(bot, db, date(2026, 9, d), "10:00") for d in (1, 2, 4)]
    octs = [await send_session(bot, db, date(2026, 10, d), "10:00") for d in (1, 2)]
    for s in sep + octs:
        db.mark(s.id, 201)
    db.mark(sep[0].id, 202)
    db.add_sick_leave(203, "2026-09-02", "2026-10-01")

    # история по месяцам
    await press(dp, bot, ADMIN, "a:sl:0", ADMIN)
    m = session.of(EditMessageText)[-1]
    assert "Октябрь 2026" in m.text and "Сентябрь 2026" in m.text
    assert buttons(m)[:2] == ["a:ml:2026-10:0", "a:ml:2026-09:0"]
    await press(dp, bot, ADMIN, "a:ml:2026-09:0", ADMIN)
    m = session.of(EditMessageText)[-1]
    assert "Занятий: 3" in m.text
    assert [b for b in buttons(m) if b.startswith("a:s:")] == [f"a:s:{s.id}" for s in sep[::-1]]

    # ведомость: все месяцы, каждый на своём листе
    await press(dp, bot, ADMIN, "a:xl:all", ADMIN)
    wb = openpyxl.load_workbook(io.BytesIO(session.of(SendDocument)[-1].document.data))
    assert wb.sheetnames == ["Сентябрь 2026", "Октябрь 2026", "Больничные"]
    ws = wb["Сентябрь 2026"]
    assert ws["A1"].value == "УФРС24-2" and ws["A2"].value == " СЕНТЯБРЬ 2026"
    assert [ws.cell(1, c).value for c in range(2, 5)] == ["Вт", "Ср", "Пт"]
    assert [ws.cell(2, c).value for c in range(2, 6)] == [1, 2, 4, "ИТОГО"]
    names = [ws.cell(r, 1).value for r in range(3, 6)]
    assert names == ["Анна Смирнова", "Иван Иванов", "Пётр Петров"]  # по алфавиту
    assert [ws.cell(3, c).value for c in range(2, 6)] == ["н", "б", "б", '=COUNTIFS(B3:D3,"Н")']
    assert [ws.cell(4, c).value for c in range(2, 5)] == [None, None, None]
    assert [ws.cell(5, c).value for c in range(2, 5)] == [None, "н", "н"]
    assert ws["C3"].fill.fgColor.rgb.endswith("FBE4D5")
    assert ws["A2"].fill.fgColor.rgb.endswith("FFFF00")
    assert ws["E2"].font.color.rgb.endswith("FF0000") and ws["A1"].font.name == "Times New Roman"


async def test_group_message_has_date(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    db.set("message_text", "Просто текст без даты")
    await send_session(bot, db, date(2026, 9, 30), "10:00")
    msg = session.of(SendMessage)[-1].text
    assert "📅 <b>Среда, 30 сентября 2026</b> · 10:00" in msg
    db.set("mode", "day")
    await send_session(bot, db, date(2026, 10, 1), "10:00")
    assert "📅 <b>Четверг, 1 октября 2026</b>\n" in session.of(SendMessage)[-1].text


def test_old_default_text_migrates():
    db = Database(":memory:")
    db.set("message_text", "📅 {день}, {дата}\n\nЕсли вы на занятии — нажмите кнопку ниже 👇")
    db._migrate()
    assert db.get("message_text") == "Если вы на занятии — нажмите кнопку ниже 👇"


# ---------- сообщество: несколько чатов и темы ----------

async def test_bind_to_forum_topic(env):
    db, bot, dp, session = env
    topic_created = Message(
        message_id=77, date=datetime.now(), chat=Chat(id=GROUP, type="supergroup", is_forum=True),
        forum_topic_created={"name": "Посещаемость", "icon_color": 0},
    )
    msg = Message(
        message_id=next(_uid), date=datetime.now(),
        chat=Chat(id=GROUP, type="supergroup", title="Сообщество ИВТ", is_forum=True),
        from_user=user(ADMIN), text="/bind", message_thread_id=77, is_topic_message=True,
        reply_to_message=topic_created,
    )
    await dp.feed_update(bot, Update(update_id=next(_uid), message=msg))
    assert db.group_chat_id == GROUP and db.group_thread_id == 77
    assert db.get("group_title") == "«Сообщество ИВТ» → тема «Посещаемость»"

    s = await send_session(bot, db, date.today(), "10:00")
    sent = session.of(SendMessage)[-1]
    assert sent.chat_id == GROUP and sent.message_thread_id == 77

    await press(dp, bot, 201, f"mark:{s.id}", GROUP, "Иван Иванов")
    assert "записано" in last_alert(session)

    # /bind в «Общем» (без темы) переключает обратно на весь чат
    await text(dp, bot, ADMIN, "/bind", GROUP, "supergroup")
    assert db.group_thread_id is None


async def test_adding_bot_to_second_chat_does_not_rebind(env):
    db, bot, dp, session = env
    me = User(id=42, is_bot=True, first_name="bot")
    await dp.feed_update(bot, Update(update_id=next(_uid),
                                     my_chat_member=member_update(GROUP, ADMIN, "left", "member", me)))
    assert db.group_chat_id == GROUP
    await dp.feed_update(bot, Update(update_id=next(_uid),
                                     my_chat_member=member_update(-2002, ADMIN, "left", "member", me)))
    assert db.group_chat_id == GROUP
    assert "по-прежнему" in [m.text for m in session.of(SendMessage) if m.chat_id == ADMIN][-1]

    # сообщения и кнопки из другого чата сообщества не влияют ни на что
    s = await send_session(bot, db, date.today(), "10:00")
    await press(dp, bot, 201, f"mark:{s.id}", -2002, "Иван Иванов")
    assert "только в чате группы" in last_alert(session)
    await text(dp, bot, 305, "привет", -2002, "supergroup")
    assert db.get_student(305) is None


# ---------- удаление занятия при ошибках Telegram ----------

@pytest.mark.parametrize("exc", ["migrate", "forbidden"])
async def test_delete_session_when_telegram_refuses(env, exc):
    from aiogram.exceptions import TelegramForbiddenError, TelegramMigrateToChat
    from aiogram.methods import DeleteMessage

    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    s = await send_session(bot, db, date.today(), "10:00")
    orig = session.make_request

    async def failing(b, m, timeout=None):
        if isinstance(m, DeleteMessage):
            if exc == "migrate":
                raise TelegramMigrateToChat(method=m, message="migrated", migrate_to_chat_id=-100777)
            raise TelegramForbiddenError(method=m, message="Forbidden: bot was kicked")
        return await orig(b, m, timeout)

    session.make_request = failing
    await press(dp, bot, ADMIN, f"a:sdy:{s.id}", ADMIN)
    assert db.get_session(s.id) is None          # занятие удалено
    alerts = [a.text or "" for a in session.of(AnswerCallbackQuery)]
    assert any("удалите вручную" in a for a in alerts)  # и кнопка не зависла


async def test_any_crash_still_answers_button(env, monkeypatch):
    db, bot, dp, session = env

    async def boom(*a, **kw):
        raise RuntimeError("boom")

    monkeypatch.setattr(admin, "show_months", boom)
    await press(dp, bot, ADMIN, "a:sl:0", ADMIN)
    assert "Что-то пошло не так" in last_alert(session)


async def test_group_migration_moves_sessions(env):
    db, bot, dp, session = env
    db.set("group_chat_id", GROUP)
    s = await send_session(bot, db, date.today(), "10:00")
    msg = Message(message_id=next(_uid), date=datetime.now(), chat=Chat(id=GROUP, type="group"),
                  from_user=user(ADMIN), migrate_to_chat_id=-100555)
    await dp.feed_update(bot, Update(update_id=next(_uid), message=msg))
    assert db.group_chat_id == -100555
    assert db.get_session(s.id).chat_id == -100555
    await press(dp, bot, 201, f"mark:{s.id}", -100555, "Иван Иванов")
    assert "записано" in last_alert(session)
